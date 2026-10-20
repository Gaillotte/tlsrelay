#!/usr/bin/env python3
# Builds docs/requirements-coverage.xlsx: every requirement of docs/SPEC.md §11, with a
# short explanation of how it is covered (HOW below) and the coverage of docs/TESTS.md §4.
# When a requirement is added or its coverage changes, update HOW and run this again.
# Requires openpyxl. The summary formulas are computed when the workbook is opened.
# Usage: tools/requirements-xlsx.py [output.xlsx]
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent

# How each requirement is covered, in a few words.
HOW = {
"GEN-01":"Built and tested in CI on Windows (MSVC) and Linux (GCC, Clang); packages installed on 5 Linux distributions; cross-platform acceptance test.",
"GEN-02":"CI job builds against OpenSSL 3.5 LTS; Windows CI checks the bundled version; startup log checked in acceptance.",
"GEN-03":"Tests show a server refuses a peer without the tlsrelay ALPN, and a client refuses a non-tlsrelay server.",
"GEN-04":"End-to-end tests send data through each of the 3 ports and verify integrity; config without ports is refused.",
"GEN-05":"Relay tests run client and server modes to one destination; config loading tests for both modes.",
"TLS-01":"Tests confirm TLS 1.2 and plaintext are refused on TLS ports; external scanners (sslyze, testssl.sh) confirm in CI.",
"TLS-02":"External TLS scanners in CI list the offered suites and groups and fail on anything not allowed.",
"TLS-03":"Tests check ALPN is negotiated and that a missing or wrong ALPN is refused.",
"TLS-04":"External TLS scan in CI confirms no session ticket or resumption is offered.",
"TLS-05":"Negative tests: server from another PKI, wrong name, wrong issuing CA, client cert used as server cert — all refused.",
"TLS-06":"Negative tests: no client certificate, wrong CA, server certificate used as client — all refused on mTLS.",
"TLS-07":"Attack tests with expired and not-yet-valid certificates on both sides; PKI validation test.",
"TLS-08":"Test closes a silent connection at the timeout; flood test shows resources stay bounded and recover.",
"TUN-01":"Unit tests for every frame type and invalid case, plus continuous fuzzing of the frame decoder.",
"TUN-02":"Tunnel tests cover open success and failure, and that early application data is buffered, not lost.",
"TUN-03":"Tunnel tests check sequence numbers and acknowledgements; out-of-order data closes with a protocol error.",
"TUN-04":"Back-pressure tests with slow application and destination: no data loss, buffers stay within limits.",
"TUN-05":"Test with a stalled destination shows the transport connection is still read (pings and acks keep flowing).",
"TUN-06":"Tests force short session lifetimes and verify renegotiation on a new connection with no data cut; acceptance R07.",
"TUN-07":"Tests check resumption tokens are single-use, bound to the old connection, and refused for another peer.",
"TUN-08":"Test lets a session exceed its lifetime plus grace and checks the server closes it with LIFETIME_EXCEEDED.",
"TUN-09":"Cut-proxy tests reset connections mid-stream; session resumes with no loss or duplication; acceptance R06.",
"TUN-10":"Tests cut the link during handshake and before OPEN_OK; tunnel reopens with no loss.",
"TUN-11":"Silent-link tests: detected after 3 keepalive intervals; resumed on TLS ports, closed on plain port.",
"TUN-12":"Test cuts a plain-port channel and verifies it is closed, not resumed.",
"TUN-13":"Test closes one direction and checks the other keeps flowing end to end.",
"TUN-14":"Tests send oversized and control-character CLOSE/OPEN_FAIL messages and check truncation and escaping.",
"CAN-01":"Tests open many parallel channels per port, each with its own session; acceptance R03 with 50 channels.",
"CAN-02":"Test opens more than max_channels and checks extra connections are refused.",
"CAN-03":"Tests check an idle channel closes at the timeout and that application traffic keeps it open.",
"CAN-04":"Tests with an unreachable destination check retries, timeout, backoff and OPEN_FAIL.",
"CAN-05":"Configuration validation tests reject inconsistent timer values.",
"WL-01":"Unit tests on whitelist parsing: unknown keys, duplicates, bad IP/CIDR/DN refused; fuzzed in CI.",
"WL-02":"Test shows a non-listed address is dropped at TCP accept, before any TLS handshake.",
"WL-03":"mTLS tests require subject and address in the same entry; DN comparison tests.",
"WL-04":"Hot-reload tests (file watch and reload command), invalid file kept out; acceptance R05.",
"WL-05":"Tests remove a peer and check its channels close, or are refused at next renegotiation per setting.",
"WL-06":"Startup test with an invalid whitelist: the service refuses to start.",
"PKI-01":"Tests generate the full CA hierarchy and check key types, constraints and purposes; acceptance R01.",
"PKI-02":"Tests create generations and check 'current' is switched atomically.",
"PKI-03":"Tests feed broken generations (missing file, key mismatch, bad chain, dates) and check refusal.",
"PKI-04":"Tests check gen-csr creates the key locally and an encrypted key is refused, never prompted; permissions checked manually (M07).",
"PKI-05":"Rotation tests under traffic: new connections use new generation, old ones drain; acceptance R08.",
"PKI-06":"Tests: invalid generation refused while running; invalid PKI blocks startup.",
"PKI-07":"Tests check the rotation guard refuses an unrelated root, and --force overrides it.",
"PKI-08":"Test runs the 3-step root change under traffic with no interruption; acceptance R09.",
"PKI-09":"Tests check pick-up by file watch, SIGHUP and reload command.",
"PKI-10":"Command-line lifecycle scripts run gen-csr, install, activate and check-pki; acceptance R01.",
"CA-01":"Tests check CA keys are encrypted, passphrase rules enforced, permissions/ACL restricted.",
"CA-02":"Signing tests refuse bad requests (no proof of possession, wrong key, missing SAN) and enforce limits.",
"CA-03":"Tests refuse a duplicate subject and mark the old certificate replaced on renewal.",
"CA-04":"Tests list issued certificates, including old-format inventories; editor inventory E06.",
"CA-05":"Tests revoke certificates and verify signed, numbered CRLs per CA; editor step E08.",
"CA-06":"Tests write the audit log, tamper with it and check detection; also run by the PKI command scripts; editor step E10.",
"CA-07":"Tests check root-change files, CRLs and commands are produced; acceptance R09.",
"CRL-01":"Tests check revoked certificates are refused at handshake on both sides, and CRL signatures verified.",
"CRL-02":"Tests reload CRLs at runtime and check revoked peers' channels close.",
"CRL-03":"Tests with invalid, stale and self-revoking CRLs check startup refusal, warnings and reports.",
"CRL-04":"Command tests for pki crl; status, check-pki and metrics report CRL state.",
"CFG-01":"Config tests check every error is reported with its path and unknown keys are refused; fuzzed.",
"CFG-02":"Same validation code used by service and editor; Qt test checks they agree; acceptance E02.",
"CFG-03":"Unit tests for durations, sizes, percentages, IPv4/IPv6 and host names.",
"CFG-04":"Round-trip tests: written files read back identically.",
"LOG-01":"Logging tests for levels, categories, text and JSON output; manual check M05.",
"LOG-02":"Tests change log level at runtime and reopen the file; acceptance R10, manual M06 on Windows.",
"LOG-03":"Tests check logs contain no secrets or data and escape peer input; acceptance R10.",
"ADM-01":"Lifecycle scripts install, start and stop real systemd and Windows services in CI; manual M01, M09.",
"ADM-02":"Lifecycle tests for graceful stop with drain, kill and restart; acceptance R12, manual M02/M03.",
"ADM-03":"Admin-channel tests for permissions and command logging; manual M04 on Windows.",
"ADM-04":"Test starts the same instance twice and checks the second is refused.",
"ADM-05":"Admin tests for status, channels, reload and check-config; acceptance R11.",
"ADM-06":"Tests reload a changed config: log levels applied, other changes flagged as needing restart.",
"ADM-07":"Service tests check the least-privilege account and capabilities; manual M07.",
"ADM-08":"Test checks metrics output follows the Prometheus text format; acceptance R11.",
"SEC-01":"Scripts check hardening flags in the Linux and Windows binaries in CI.",
"SEC-02":"Sanitizer builds (ASan, UBSan, TSan), fuzzing and static analysis in CI; longer fuzzing before production.",
"SEC-03":"12 attack tests in CI (fake server, hostile frames, floods, log injection); manual penetration test.",
"SEC-04":"Dependency check script in CI flags known-vulnerable versions.",
"PERF-01":"CI runs 1,000 concurrent channels; 5,000 measured manually on the performance machine.",
"PERF-02":"Performance script in CI measures per-channel throughput; acceptance R04.",
"PERF-03":"Performance script in CI measures added latency; acceptance R04.",
"PERF-04":"CI runs a 4-minute soak test; full 72 h run done manually before production.",
"PKG-01":"CI builds and installs/uninstalls .deb, .rpm and .msi packages; in-place upgrade checked manually.",
"PKG-02":"Package tests check licence and notice files are shipped; acceptance E01.",
"PKG-03":"CI release job publishes packages with their SHA-256 hashes; manual check M08.",
"EDT-01":"Qt tests check editor validation matches the service and highlights faulty fields; acceptance E02.",
"EDT-02":"Qt tests: save refused while invalid, diff preview, .bak kept, atomic write; acceptance E03/E04.",
"EDT-03":"Qt tests export Linux and Windows deployment packages and check contents; acceptance E05.",
"EDT-04":"Qt tests run the full PKI panel cycle (sign, renew, revoke, CRL, audit); acceptance E06–E09.",
"EDT-05":"Qt tests run the root-change wizard and check the files it prepares; acceptance E10.",
"EDT-06":"MSI install tested in CI with the editor; headless --check tested; acceptance E01.",
}

STATUS_FILL = {"Auto": "E2EFDA", "Auto + manual": "FFF2CC", "Partial (manual)": "FCE4D6"}
FONT = "Arial"
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
BODY_FONT = Font(name=FONT, size=10)
NOTE_FONT = Font(name=FONT, italic=True, size=9, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def cells(line):
    """Cells of a Markdown table row (escaped pipes kept)."""
    return [c.strip() for c in re.split(r"(?<!\\)\|", line.strip())[1:-1]]


def plain(text):
    """Markdown cell text without code quotes; an em dash alone means none."""
    text = re.sub(r"`([^`]*)`", r"\1", text).replace("\\|", "|").strip()
    return "None" if text == "—" else text


def section(path, start, end):
    return (ROOT / "docs" / path).read_text(encoding="utf-8").split(start)[1].split(end)[0]


def read_requirements():
    reqs, group = [], None
    for line in section("SPEC.md", "## 11. Requirements", "## 12.").splitlines():
        m = re.match(r"### 11\.\d+ (.*)", line)
        if m:
            group = m.group(1)
        elif re.match(r"\| [A-Z]+-\d+ ", line):
            c = cells(line)
            reqs.append({"id": c[0], "group": group, "text": c[1], "ref": c[2], "release": c[3]})
    coverage = {}
    for line in section("TESTS.md", "## 4. Requirements coverage", "## 5.").splitlines():
        if re.match(r"\| [A-Z]+-\d+ ", line):
            c = cells(line)
            coverage[c[0]] = {"auto": c[2], "manual": c[3], "status": c[4]}
    ids = {r["id"] for r in reqs}
    if ids != set(coverage) or ids != set(HOW):
        sys.exit("requirement lists differ: SPEC.md §11 only: %s; TESTS.md §4 only: %s; HOW only: %s; not in HOW: %s"
                 % (sorted(ids - set(coverage)), sorted(set(coverage) - ids), sorted(set(HOW) - ids),
                    sorted(ids - set(HOW))))
    for r in reqs:
        r.update(coverage[r["id"]])
    return reqs


def header(ws, row, names):
    for col, name in enumerate(names, 1):
        c = ws.cell(row, col, name)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")


def requirements_sheet(ws, reqs):
    ws.title = "Requirements"
    ws["A1"] = "tlsrelay — Requirements and how they are covered"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = ("Source: docs/SPEC.md §11 (requirements) and docs/TESTS.md §4 (coverage). "
                "Releases: 1.0 = relay service, 1.1 = tlsrelay-admin editor.")
    ws["A2"].font = NOTE_FONT
    columns = [("ID", 10), ("Group", 22), ("Requirement", 60), ("Release", 9), ("Spec ref.", 12),
               ("How it is covered (in a few words)", 60), ("Coverage status", 16),
               ("Automated tests (detail)", 60), ("Acceptance / manual checks", 26)]
    first = 4
    header(ws, first, [name for name, _ in columns])
    for col, (_, width) in enumerate(columns, 1):
        ws.column_dimensions[chr(64 + col)].width = width
    for row, r in enumerate(reqs, first + 1):
        values = [r["id"], r["group"], plain(r["text"]), r["release"], plain(r["ref"]), HOW[r["id"]],
                  r["status"], plain(r["auto"]), plain(r["manual"])]
        for col, value in enumerate(values, 1):
            c = ws.cell(row, col, value)
            c.font, c.alignment, c.border = BODY_FONT, WRAP, BORDER
        ws.cell(row, 1).font = Font(name=FONT, size=10, bold=True)
        ws.cell(row, 7).fill = PatternFill("solid", fgColor=STATUS_FILL[r["status"]])
    last = first + len(reqs)
    ws.row_dimensions[first].height = 30
    ws.freeze_panes = ws.cell(first + 1, 2)
    ws.auto_filter.ref = f"A{first}:I{last}"
    ws.print_title_rows = f"{first}:{first}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return first + 1, last


def summary_sheet(ws, reqs, first, last):
    ws["A1"] = "Coverage summary"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = "Counts are formulas over the Requirements sheet."
    ws["A2"].font = NOTE_FONT
    header(ws, 4, ["Group", "Requirements", "Auto", "Auto + manual", "Partial (manual)", "% fully automated"])
    groups = f"Requirements!$B${first}:$B${last}"
    statuses = f"Requirements!$G${first}:$G${last}"
    row = 5
    for group in dict.fromkeys(r["group"] for r in reqs):
        ws.cell(row, 1, group)
        ws.cell(row, 2, f"=COUNTIF({groups},A{row})")
        for col, status in ((3, "Auto"), (4, "Auto + manual"), (5, "Partial (manual)")):
            ws.cell(row, col, f'=COUNTIFS({groups},$A{row},{statuses},"{status}")')
        ws.cell(row, 6, f"=IF(B{row}=0,0,C{row}/B{row})")
        row += 1
    total = row
    ws.cell(total, 1, "Total")
    for col in "BCDE":
        ws[f"{col}{total}"] = f"=SUM({col}5:{col}{total - 1})"
    ws[f"F{total}"] = f"=IF(B{total}=0,0,C{total}/B{total})"
    ws[f"A{total + 1}"] = "Automated (fully or with a manual complement)"
    ws[f"F{total + 1}"] = f"=IF(B{total}=0,0,(C{total}+D{total})/B{total})"
    for row in range(5, total + 2):
        for col in range(1, 7):
            c = ws.cell(row, col)
            c.font, c.border = Font(name=FONT, size=10, bold=row >= total), BORDER
            if col == 6:
                c.number_format = "0.0%"
    ws.column_dimensions["A"].width = 40
    for col in "BCDEF":
        ws.column_dimensions[col].width = 15
    ws.freeze_panes = "A5"


def legend_sheet(ws):
    rows = [
        ("Coverage status", ""),
        ("Auto", "Fully verified by automated tests on every push (GitHub Actions CI)."),
        ("Auto + manual", "Automated in CI, with a manual complement during acceptance or before production."),
        ("Partial (manual)", "Automated at reduced scale in CI; full verification is manual "
                             "(performance machine, long runs)."),
        ("", ""),
        ("Test references", ""),
        ("GoogleTest names (e.g. TlsTest.ServerRejectsTls12)",
         "Unit and integration tests in tests/*.cpp, run by ctest on Linux and Windows."),
        ("Qt test names (e.g. pki_panel_full_cycle)", "Editor tests in tests/admin/test_admin.cpp."),
        ("Scripts (systemd.sh, tls-scan.sh, perf.sh …)",
         "Lifecycle, security and performance scripts in tests/lifecycle, tests/perf and tools/."),
        ("R01–R13", "Automated acceptance scenarios (tools/acceptance.ps1), docs/ACCEPTANCE.md."),
        ("M01–M09", "Manual acceptance steps for the service, docs/ACCEPTANCE.md."),
        ("E01–E10", "Acceptance steps for the editor, docs/ACCEPTANCE.md."),
        ("Spec ref.", "Section of docs/SPEC.md (§n), PLAN.md or EDITOR.md describing the requirement."),
    ]
    for row, (term, meaning) in enumerate(rows, 1):
        ws.cell(row, 1, term).font = Font(name=FONT, size=10, bold=not meaning or term in STATUS_FILL)
        ws.cell(row, 2, meaning).font = BODY_FONT
        ws.cell(row, 2).alignment = WRAP
        if term in STATUS_FILL:
            ws.cell(row, 1).fill = PatternFill("solid", fgColor=STATUS_FILL[term])
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 90


def main():
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "requirements-coverage.xlsx"
    reqs = read_requirements()
    wb = Workbook()
    first, last = requirements_sheet(wb.active, reqs)
    summary_sheet(wb.create_sheet("Summary"), reqs, first, last)
    legend_sheet(wb.create_sheet("Legend"))
    wb.save(output)
    print(f"{output}: {len(reqs)} requirements")


if __name__ == "__main__":
    main()
